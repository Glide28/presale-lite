"""Извлечение текста из загруженных документов (txt, md, docx, pdf)."""
from __future__ import annotations

import io
from pathlib import PurePath

ALLOWED_EXTENSIONS = {".txt", ".md", ".docx", ".pdf"}
MAX_TEXT_CHARS = 60_000  # ограничиваем контекст, передаваемый в LLM


class DocumentError(ValueError):
    """Файл не прошёл проверку или не удалось извлечь текст."""


def extract_text(filename: str, content: bytes, max_bytes: int = 5 * 1024 * 1024) -> str:
    ext = PurePath(filename or "").suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        allowed = ", ".join(sorted(ALLOWED_EXTENSIONS))
        raise DocumentError(f"Неподдерживаемый формат '{ext or filename}'. Допустимо: {allowed}")
    if not content:
        raise DocumentError("Файл пустой")
    if len(content) > max_bytes:
        raise DocumentError(f"Файл больше {max_bytes // (1024 * 1024)} МБ")

    if ext in {".txt", ".md"}:
        text = _decode(content)
    elif ext == ".docx":
        text = _read_docx(content)
    else:
        text = _read_pdf(content)

    text = text.strip()
    if not text:
        raise DocumentError("В документе не найден текст")
    return text[:MAX_TEXT_CHARS]


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _read_docx(content: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise DocumentError("Для .docx установите python-docx") from exc
    try:
        doc = Document(io.BytesIO(content))
    except Exception as exc:
        raise DocumentError("Не удалось прочитать .docx") from exc
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _read_pdf(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise DocumentError("Для .pdf установите pypdf") from exc
    try:
        reader = PdfReader(io.BytesIO(content))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception as exc:
        raise DocumentError("Не удалось прочитать .pdf") from exc
